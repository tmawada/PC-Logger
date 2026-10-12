from gevent import monkey
monkey.patch_all()
from flask import Flask, render_template
from config import Config
from database import init_db, db
from routes import api
from models import PcLog
from datetime import datetime, timedelta
import threading
import time
import pytz
import sys
import os
import traceback
import io
from flask import Response
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from gevent.pywsgi import WSGIServer
from prometheus_flask_exporter import PrometheusMetrics
from prometheus_client import Gauge


sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "model"))
from seed import ROOM_RATES, get_room_rate, seed_data  # noqa: E402


rooms =  ['601','602','603','604','605','606','608','610','611',
        '621','622','623','624','625','626','627','628','629','630'
        ,'706', '708', '710', '722', '724', '721', '723', '725'
        , '727', '729', '730', '731']

pcs_per_room = 41
events = {'boot', 'restart'}


GRAFANA_URL = os.getenv("GRAFANA_URL")

lab_pcs_status = Gauge(
    'lab_pcs_status', 'Current PC status count per room', ['room', 'status']
)

pc_lab_top_longest_running_seconds = Gauge(
    'pc_lab_top_longest_running_seconds',
    'Uptime in seconds of the top 20 longest-running PCs',
    ['rank', 'room', 'hostname', 'pcnum']
)


pc_lab_cost_per_room = Gauge(
    'pc_lab_cost_per_room',
    'Live electricity cost (Rp/hour) for the room, based on PCs currently on',
    ['rank', 'room']
)

pc_lab_kwh_per_room = Gauge(
    'pc_lab_kwh_per_room',
    'Live power draw (kWh/hour) for the room, based on PCs currently on',
    ['rank', 'room']
)


def create_app():
    app = Flask(__name__)
    app.config.from_object(Config)
    metrics = PrometheusMetrics(app)
    try:
        init_db(app)
    except Exception:
        traceback.print_exc()
    app.register_blueprint(api)

    @app.route("/main")
    def main_menu():
        pcs = PcLog.query.order_by(PcLog.id.desc()).limit(3000).all()
        pcs = [p.to_dict() for p in pcs]
        rooms_data = build_room_status(pcs, rooms)
        generated_at = datetime.now().strftime('%Y-%m-%d %H:%M:%S')

        total_on = 0
        total_cost_per_hour = 0.0
        total_kwh_per_hour = 0.0
        total_cost_today = 0.0
        total_kwh_today = 0.0
        total_on_hours_today = 0.0
        for room in rooms:
            stats = rooms_data.get(room, {'on': 0})
            rate = get_room_rate(room)
            room_cost_per_hour = rate['cost_per_pc_hour'] * stats['on']
            total_on += stats['on']
            total_cost_per_hour += room_cost_per_hour
            total_kwh_per_hour += rate['kwh_per_pc_hour'] * stats['on']

            room_prefix = f"BN{room}%"
            room_logs = PcLog.query.filter(PcLog.hostname.like(room_prefix)).order_by(PcLog.id.desc()).limit(5000).all()
            room_logs = [d.to_dict() for d in room_logs]
            today_on_seconds = compute_today_on_seconds(room_logs)
            room_on_hours_today = sum(today_on_seconds.values()) / 3600.0
            total_kwh_today += rate['kwh_per_pc_hour'] * room_on_hours_today
            total_cost_today += rate['cost_per_pc_hour'] * room_on_hours_today
            total_on_hours_today += room_on_hours_today

        return render_template(
            "mainmenu.html",
            rooms=rooms,
            rooms_data=rooms_data,
            pcs_per_room=pcs_per_room,

            generated_at=generated_at,
            total_on=total_on,
            total_cost_per_hour=round(total_cost_per_hour),
            total_kwh_per_hour=round(total_kwh_per_hour, 3),
            total_cost_today=round(total_cost_today),
            total_kwh_today=round(total_kwh_today, 3),
            total_on_hours_today=round(total_on_hours_today, 2),
        )

    @app.route("/main/export")
    def export_main():
        pcs = PcLog.query.order_by(PcLog.id.desc()).limit(3000).all()
        pcs = [p.to_dict() for p in pcs]
        rooms_data = build_room_status(pcs, rooms)

        rows = []
        for room in rooms:
            stats = rooms_data.get(room, {'on': 0, 'off': 0, 'unknown': pcs_per_room})
            rate = get_room_rate(room)

            room_prefix = f"BN{room}%"
            room_logs = PcLog.query.filter(PcLog.hostname.like(room_prefix)).order_by(PcLog.id.desc()).limit(5000).all()
            room_logs = [d.to_dict() for d in room_logs]
            today_on_seconds = compute_today_on_seconds(room_logs)
            room_on_hours_today = sum(today_on_seconds.values()) / 3600.0

            rows.append({
                "Room": room,
                "PCs On": stats['on'],
                "PCs Off": stats['off'],
                "Unknown": stats['unknown'],
                "Rp/hour (live)": round(rate['cost_per_pc_hour'] * stats['on']),
                "kWh/hour (live)": round(rate['kwh_per_pc_hour'] * stats['on'], 3),
                "On-hours today (sum, all PCs)": round(room_on_hours_today, 2),
                "kWh today": round(rate['kwh_per_pc_hour'] * room_on_hours_today, 3),
                "Rp today": round(rate['cost_per_pc_hour'] * room_on_hours_today),
            })

        return build_excel_response(
            sheet_title="All Rooms",
            headers=["Room", "PCs On", "PCs Off", "Unknown", "Rp/hour (live)",
                     "kWh/hour (live)", "On-hours today (sum, all PCs)", "kWh today", "Rp today"],
            rows=rows,
            filename=f"pc-logger-rooms-{datetime.now().strftime('%Y%m%d-%H%M%S')}.xlsx",
        )

    # def query_room_number(hostname):
        

    @app.route("/main/room/<room_id>")
    def get_room_details(room_id):
        room_prefix = f"BN{room_id}%"
        room_id_data = PcLog.query.filter(PcLog.hostname.like(room_prefix)).order_by(PcLog.id.desc()).limit(5000).all()
        room_id_data = [d.to_dict() for d in room_id_data]
        # print("Room id data : ", room_id_data, flush=True)
        pcs = build_pc_details(room_id_data)
        on_count = sum(1 for p in pcs if p['status'] == 'on')
        off_count = sum(1 for p in pcs if p['status'] == 'off')
        rate = get_room_rate(room_id)

        cost_per_hour = round(rate['cost_per_pc_hour'] * on_count)

        today_on_seconds = compute_today_on_seconds(room_id_data)
        on_hours_today = sum(today_on_seconds.values()) / 3600.0
        kwh_today = round(rate['kwh_per_pc_hour'] * on_hours_today, 3)
        cost_today = round(rate['cost_per_pc_hour'] * on_hours_today)

        return render_template(
            "room_details.html",
            room_id=room_id,
            pcs=pcs,
            on_count=on_count,
            off_count=off_count,
            unknown_count=pcs_per_room - on_count - off_count,
            kwh_per_hour=round(rate['kwh_per_pc_hour'] * on_count, 3),
            cost_per_hour=cost_per_hour,
            cost_per_pc_hour=rate['cost_per_pc_hour'],
            kwh_per_pc_hour=rate['kwh_per_pc_hour'],
            on_hours_today=round(on_hours_today, 2),
            kwh_today=kwh_today,
            cost_today=cost_today,
        )

    @app.route("/main/room/<room_id>/export")
    def export_room(room_id):
        room_prefix = f"BN{room_id}%"
        room_id_data = PcLog.query.filter(PcLog.hostname.like(room_prefix)).order_by(PcLog.id.desc()).limit(5000).all()
        room_id_data = [d.to_dict() for d in room_id_data]
        pcs = build_pc_details(room_id_data)
        rate = get_room_rate(room_id)
        today_on_seconds = compute_today_on_seconds(room_id_data)

        rows = []
        for pc in pcs:
            hostname = f"BN{room_id}{pc['pcnum']:02d}"
            on_hours_today = today_on_seconds.get(hostname, 0.0) / 3600.0
            rows.append({
                "PC": pc['pcnum'],
                "Status": pc['status'],
                "Last boot": pc['last_boot'],
                "Last restart": pc['last_restart'],
                "Last shutdown": pc['last_off'],
                "Uptime (current session)": pc['uptime'] or '',
                "On-hours today": round(on_hours_today, 2),
                "kWh today": round(rate['kwh_per_pc_hour'] * on_hours_today, 3),
                "Rp today": round(rate['cost_per_pc_hour'] * on_hours_today),
            })

        return build_excel_response(
            sheet_title=f"Room {room_id}",
            headers=["PC", "Status", "Last boot", "Last restart", "Last shutdown",
                     "Uptime (current session)", "On-hours today", "kWh today", "Rp today"],
            rows=rows,
            filename=f"pc-logger-room-{room_id}-{datetime.now().strftime('%Y%m%d-%H%M%S')}.xlsx",
        )


    @app.route("/main/room/<room_id>/export/<start_date>to<end_date>")
    def export_room_ranged(room_id, start_date, end_date):
        room_prefix = f"BN{room_id}%"

        room_id_data = PcLog.query.filter(PcLog.hostname.like(room_prefix)).filter(PcLog.received_at.between(start_date, end_date)).order_by(PcLog.id.desc()).limit(5000).all()

        room_id_data = [d.to_dict() for d in room_id_data]
        pcs = build_pc_details(room_id_data, range_start=start_date, range_end=end_date)
        rate = get_room_rate(room_id)
        range_on_seconds = compute_on_seconds_in_range(room_id_data, start_date, end_date)

        rows = []
        for pc in pcs:
            hostname = f"BN{room_id}{pc['pcnum']:02d}"
            on_hours_range = range_on_seconds.get(hostname, 0.0) / 3600.0
            rows.append({
                "PC": pc['pcnum'],
                "Status": pc['status'],
                "Last boot": pc['last_boot'],
                "Last restart": pc['last_restart'],
                "Last shutdown": pc['last_off'],
                "Total uptime in range": pc['total_uptime'] or '',
                "On-hours in range": round(on_hours_range, 2),
                "kWh in range": round(rate['kwh_per_pc_hour'] * on_hours_range, 3),
                "Rp in range": round(rate['cost_per_pc_hour'] * on_hours_range),
            })

        return build_excel_response(
            sheet_title=f"Room {room_id}",
            headers=["PC", "Status", "Last boot", "Last restart", "Last shutdown",
                    "Total uptime in range", "On-hours in range", "kWh in range", "Rp in range"],
            rows=rows,
            filename=f"pc-logger-room-{room_id}-{start_date}-{end_date}.xlsx",
        )

    @app.route("/main/export/<start_date>to<end_date>")
    def export_main_ranged(start_date, end_date):
        pcs = PcLog.query.order_by(PcLog.id.desc()).limit(3000).all()
        pcs = [p.to_dict() for p in pcs]
        rooms_data = build_room_status(pcs, rooms)
        total_cost = 0
        rows = []
        for room in rooms:
            stats = rooms_data.get(room, {'on': 0, 'off': 0, 'unknown': pcs_per_room})
            rate = get_room_rate(room)

            room_prefix = f"BN{room}%"
            room_logs = PcLog.query.filter(PcLog.hostname.like(room_prefix)).filter(PcLog.received_at.between(start_date, end_date)).order_by(PcLog.id.desc()).limit(5000).all()
            
            room_logs = [d.to_dict() for d in room_logs]
            today_on_seconds = compute_today_on_seconds(room_logs)
            room_on_hours_today = sum(today_on_seconds.values()) / 3600.0
            cost_today = round(rate['cost_per_pc_hour'] * room_on_hours_today)
            total_cost = cost_today + total_cost
            rows.append({
                "Room": room,
                "Rp/hour": round(rate['cost_per_pc_hour'] * stats['on']),
                "kWh/hour": round(rate['kwh_per_pc_hour'] * stats['on'], 3),
                # "On-hours total": round(room_on_hours_today, 2),
                "kWh total": round(rate['kwh_per_pc_hour'] * room_on_hours_today, 3),
                "Rp total": round(rate['cost_per_pc_hour'] * room_on_hours_today),
            })
        rows.append({
            "Room": "TOTAL",
            "Rp total": total_cost,
        })
        return build_excel_response(
            sheet_title="All Rooms",
            headers=["Room","Rp/hour)",
                     "kWh/hour", 
                     # "On-hours today (sum, all PCs)", 
                     "kWh total", 
                     "Rp total"],
            rows=rows,
            filename=f"pc-logger-rooms-{datetime.now().strftime('%Y%m%d-%H%M%S')}.xlsx",
        )

    @app.route("/monitoring")
    def monitoring():
        return render_template("monitoring.html", grafana_url=GRAFANA_URL)

    return app

app = create_app()
with app.app_context():
    db.create_all()
    try:
        seed_data()
    except Exception:
        import traceback
        traceback.print_exc()

def build_excel_response(sheet_title, headers, rows, filename):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31] if sheet_title else "Sheet1"

    header_font = Font(bold=True)
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(header))

    for col_idx, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 15

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return Response(
        buffer.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )

def format_date(value):
    if isinstance(value, str):
        try:
            value = datetime.fromisoformat(value)
        except ValueError:
            return value

    if isinstance(value, datetime):
        return value.strftime('%b %d, %H:%M')

    return value

def parse_datetime(value):
    if isinstance(value, datetime):
        return value
    if not value:
        return None
    try:
        return datetime.fromisoformat(value)
    except (TypeError, ValueError):
        return None

def format_duration(time):
    total_minutes = int(time.total_seconds() // 60)

    if total_minutes <= 0:
        return 'from the future'
    if total_minutes < 1:
        return 'less than a minute'
    days, rem = divmod(total_minutes, 1440)
    hours, minutes = divmod(rem, 60)
    parts = []
    if days:
        parts.append(f'{days}d')
    if hours:
        parts.append(f'{hours}h')
    if minutes and not days:
        parts.append(f'{minutes}m')
    return ' '.join(parts) if parts else '0m'


def update_longest_running(top_n=20):
    logs = PcLog.query.order_by(PcLog.id.desc()).all()
    logs = [p.to_dict() for p in logs]

    per_host = {}
    for entry in logs:
        hostname = entry.get('hostname', '')
        if len(hostname) < 5 or not hostname[-5:].isdigit():
            continue
        room = hostname[-5:-2]
        event_type = entry.get('event_type')
        created_at = parse_datetime(entry.get('created_at'))
        if created_at is None:
            continue
        pcnum = int(hostname[-2:])
        host = per_host.setdefault(hostname, {'room': room,'pcnum': pcnum, 'last_boot': None, 'last_off': None})
        if event_type == "Shutdown":
            if host['last_off'] is None or created_at > host['last_off']:
                host['last_off'] = created_at
        else:
            if host['last_boot'] is None or created_at > host['last_boot']:
                host['last_boot'] = created_at

    indonesia_timezone = pytz.timezone('Asia/Jakarta')
    time_raw =  datetime.now(indonesia_timezone)
    now = time_raw.replace(tzinfo=None)
    running = []
    for hostname, host in per_host.items():
        is_on = host['last_boot'] and (not host['last_off'] or host['last_boot'] > host['last_off'])
        if not is_on:
            continue
        uptime_seconds = (now - host['last_boot']).total_seconds()
        running.append((uptime_seconds, hostname, host['room'], host['pcnum']))

    running.sort(reverse=True)
    top = running[:top_n]

    pc_lab_top_longest_running_seconds.clear()
    for rank, (uptime_seconds, hostname, room, pcnum) in enumerate(top, start=1):
        pc_lab_top_longest_running_seconds.labels(rank=str(rank), room=room, hostname=f'PC {hostname}', pcnum=pcnum).set(uptime_seconds)



def update_metrics_once(top_n=20):
    pcs = PcLog.query.order_by(PcLog.id.desc()).limit(3000).all()
    pcs = [p.to_dict() for p in pcs]
    rooms_data = build_room_status(pcs, rooms)

    for room, stats in rooms_data.items():
        lab_pcs_status.labels(room=room, status='on').set(stats['on'])
        lab_pcs_status.labels(room=room, status='off').set(stats['off'])
        lab_pcs_status.labels(room=room, status='unknown').set(stats['unknown'])
    
    room_costs = []
    for room, stats in rooms_data.items():
        rate = get_room_rate(room)
        cost_per_hour = rate['cost_per_pc_hour'] * stats['on']
        kwh_per_hour = rate['kwh_per_pc_hour'] * stats['on']

        room_costs.append((cost_per_hour, kwh_per_hour, room))

    room_costs.sort(reverse=True)
    top = room_costs[:top_n]

    pc_lab_cost_per_room.clear()
    pc_lab_kwh_per_room.clear()
    for rank, (cost_per_hour, kwh_per_hour, room) in enumerate(top, start=1):
        pc_lab_cost_per_room.labels(rank=str(rank), room=room).set(round(cost_per_hour, 2))
        pc_lab_kwh_per_room.labels(rank=str(rank), room=room).set(round(kwh_per_hour, 3))





def build_room_status(pcs, rooms):
    latest_by_host = {}
    for p in pcs:
        hostname = p.get('hostname', '')
        if hostname in latest_by_host:
            continue
        if len(hostname) < 5 or not hostname[-5:].isdigit():
            continue
        room = hostname[-5:-2]
        pcnum = int(hostname[-2:])
        if room not in rooms or not (1 <= pcnum <= pcs_per_room):
            continue
        if p.get('event_type') == "Shutdown":
            pcStatus = 'off'
        else:
            pcStatus = 'on'
        latest_by_host[hostname] = {
            'room': room,
            'status': pcStatus,
        }
    data = {}
    for room in rooms:
        on_count = sum(1 for e in latest_by_host.values() if e['room'] == room and e['status'] == 'on')
        off_count = sum(1 for e in latest_by_host.values() if e['room'] == room and e['status'] == 'off')
        data[room] = {
            'on': on_count,
            'off': off_count,
            'unknown': pcs_per_room - on_count - off_count,
        }
    return data


def compute_today_on_seconds(logs):
    indonesia_timezone = pytz.timezone('Asia/Jakarta')
    now = datetime.now(indonesia_timezone).replace(tzinfo=None)
    midnight = now.replace(hour=0, minute=0, second=0, microsecond=0)

    per_host_events = {}
    per_host_last_before_midnight = {}

    for entry in logs:
        hostname = entry.get('hostname', '')
        if len(hostname) < 5 or not hostname[-5:].isdigit():
            continue
        event_type = entry.get('event_type')
        created_at = parse_datetime(entry.get('created_at'))
        if created_at is None:
            continue

        if created_at < midnight:
            prev = per_host_last_before_midnight.get(hostname)
            if prev is None or created_at > prev[0]:
                per_host_last_before_midnight[hostname] = (created_at, event_type)
            continue

        if created_at > now:
            continue

        per_host_events.setdefault(hostname, []).append((created_at, event_type))

    on_seconds = {}
    hosts = set(per_host_events) | set(per_host_last_before_midnight)
    for hostname in hosts:
        events_today = sorted(per_host_events.get(hostname, []), key=lambda e: e[0])

        last_before = per_host_last_before_midnight.get(hostname)
        is_on = bool(last_before) and last_before[1] != "Shutdown"

        cursor = midnight
        accumulated = 0.0
        for created_at, event_type in events_today:
            if is_on and created_at > cursor:
                accumulated += (created_at - cursor).total_seconds()
            cursor = created_at
            is_on = event_type != "Shutdown"

        if is_on and now > cursor:
            accumulated += (now - cursor).total_seconds()

        if accumulated > 0:
            on_seconds[hostname] = accumulated

    return on_seconds


def compute_on_seconds_in_range(logs, range_start, range_end):
    range_start = parse_datetime(range_start)
    range_end = parse_datetime(range_end)
    if range_start is None or range_end is None:
        return {}
    if range_end < range_start:
        range_start, range_end = range_end, range_start

    per_host_events = {}
    per_host_last_before_start = {}

    for entry in logs:
        hostname = entry.get('hostname', '')
        if len(hostname) < 5 or not hostname[-5:].isdigit():
            continue
        event_type = entry.get('event_type')
        created_at = parse_datetime(entry.get('created_at'))
        if created_at is None:
            continue

        if created_at < range_start:
            prev = per_host_last_before_start.get(hostname)
            if prev is None or created_at > prev[0]:
                per_host_last_before_start[hostname] = (created_at, event_type)
            continue

        if created_at > range_end:
            continue

        per_host_events.setdefault(hostname, []).append((created_at, event_type))

    on_seconds = {}
    hosts = set(per_host_events) | set(per_host_last_before_start)
    for hostname in hosts:
        events_in_range = sorted(per_host_events.get(hostname, []), key=lambda e: e[0])

        last_before = per_host_last_before_start.get(hostname)
        is_on = bool(last_before) and last_before[1] != "Shutdown"

        cursor = range_start
        accumulated = 0.0
        for created_at, event_type in events_in_range:
            if is_on and created_at > cursor:
                accumulated += (created_at - cursor).total_seconds()
            cursor = created_at
            is_on = event_type != "Shutdown"

        if is_on and range_end > cursor:
            accumulated += (range_end - cursor).total_seconds()

        if accumulated > 0:
            on_seconds[hostname] = accumulated

    return on_seconds


def build_pc_details(logs, range_start=None, range_end=None):
    per_host = {}
    for entry in logs:
        hostname = entry.get('hostname', '')
        if len(hostname) < 5 or not hostname[-5:].isdigit():
            continue
        pcnum = int(hostname[-2:])
        if not (1 <= pcnum <= pcs_per_room):
            continue

        event_type = entry.get('event_type')
        created_at = parse_datetime(entry.get('created_at'))
        if created_at is None:
            continue

        host = per_host.setdefault(hostname, {
            'pcnum': pcnum, 'last_boot': None, 'last_restart': None, 'last_off': None
        })

        if event_type == "Shutdown":
            if host['last_off'] is None or created_at > host['last_off']:
                host['last_off'] = created_at
        elif event_type == "Restart":
            if host['last_restart'] is None or created_at > host['last_restart']:
                host['last_restart'] = created_at
        else:
            if host['last_boot'] is None or created_at > host['last_boot']:
                host['last_boot'] = created_at

    indonesia_timezone = pytz.timezone('Asia/Jakarta')
    time_raw =  datetime.now(indonesia_timezone)
    now = time_raw.replace(tzinfo=None)

    range_on_seconds = {}
    if range_start is not None and range_end is not None:
        range_on_seconds = compute_on_seconds_in_range(logs, range_start, range_end)

    by_pcnum = {}
    for hostname, host in per_host.items():
        last_boot = host['last_boot']
        last_restart = host['last_restart']
        last_off = host['last_off']
        last_on = max((t for t in (last_boot, last_restart) if t is not None), default=None)

        if last_on and (not last_off or last_on > last_off):
            status = 'on'
        elif last_off:
            status = 'off'
        else:
            status = 'unknown'

        uptime = None
        uptime_hours = None
        if status == 'on' and last_on:
            _delta = now - last_on
            uptime = format_duration(_delta)
            uptime_hours = _delta.total_seconds() / 3600
        elif status == 'off' and last_on and last_off:
            _delta = last_off - last_on
            uptime = format_duration(_delta)
            uptime_hours = _delta.total_seconds() / 3600
        if uptime_hours is not None and uptime_hours < 0:
            uptime_hours = None

        if status == 'on':
            label = f'Turned on at {format_date(last_on)}'
        elif status == 'off':
            label = f'Shut down at {format_date(last_off)}'
        else:
            label = 'No recent log'

        total_seconds = range_on_seconds.get(hostname)
        total_uptime = None
        total_uptime_hours = None
        if total_seconds is not None:
            total_uptime_hours = total_seconds / 3600
            total_uptime = format_duration(timedelta(seconds=total_seconds))

        by_pcnum[host['pcnum']] = {
            'pcnum': host['pcnum'],
            'status': status,
            'label': label,
            'raw_last_boot': last_boot,
            'raw_last_restart': last_restart,
            'raw_last_off': last_off,
            'last_boot': format_date(last_boot) or 'No boot log',
            'last_restart': format_date(last_restart) or 'No restart log',
            'last_off': format_date(last_off) or 'No shutdown log',
            'uptime': uptime,
            'uptime_hours': uptime_hours,
            'total_uptime': total_uptime,
            'total_uptime_hours': total_uptime_hours,
        }

    pcs = []
    for i in range(1, pcs_per_room + 1):
        pcs.append(by_pcnum.get(i) or {
            'pcnum': i,
            'status': 'unknown',
            'label': 'No recent log',
            'last_boot': 'No boot log',
            'last_restart': 'No restart log',
            'last_off': 'No shutdown log',
            'uptime': None,
            'uptime_hours': None,
            'total_uptime': None,
            'total_uptime_hours': None,
        })
    return pcs



def update_metrics_loop():
    while True:
        try:
            with app.app_context():
                update_metrics_once()
                update_longest_running()
        except Exception as e:
                import traceback
                traceback.print_exc()
        time.sleep(30)


threading.Thread(target=update_metrics_loop, daemon=True).start()

if __name__ == "__main__":
    http_server = WSGIServer(('0.0.0.0', 5000), app)
    http_server.serve_forever()


